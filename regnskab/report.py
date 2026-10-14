"""Samler bogføringslinjer og skriver Excel-rapporten."""

from collections import defaultdict
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import classify as C

DK = ZoneInfo("Europe/Copenhagen")
L = C.LAMPORTS


def berig(linjer, kurser, navne):
    """Tilføjer dato, SOL-beløb og DKK til hver linje."""
    manglende_kurs = set()
    for l in linjer:
        utc = datetime.fromtimestamp(l["ts"], tz=timezone.utc)
        l["tid_dk"] = utc.astimezone(DK).replace(tzinfo=None)
        l["dato_dk"] = l["tid_dk"].date()
        k = kurser.get(utc.date())
        if k is None:
            manglende_kurs.add(utc.date())
        l["sol_dkk_kurs"] = k[2] if k else None
        for felt in ("sol", "netvaerksgebyr", "jito", "rent", "proxima_gebyr"):
            l[felt + "_sol"] = l[felt] / L
            l[felt + "_dkk"] = l[felt] / L * k[2] if k else None
        l["token_ui"] = l["token"] / (10 ** l["decimals"]) if l["decimals"] else l["token"]
        l["navn"] = navne.get(l["mint"], "") if l["mint"] else ""
        l["wallet_navn"] = navne.get(l["wallet"], "")
    return manglende_kurs


def _sum(rows, felt):
    v = [r[felt] for r in rows if r[felt] is not None]
    return sum(v) if v else 0


def lav_rapport(alle_linjer, fra, til, oprettet, kurser, navne, cfg, filnavn):
    manglende = berig(alle_linjer, kurser, navne)
    periode = [l for l in alle_linjer if fra <= l["dato_dk"] <= til]
    periode.sort(key=lambda l: (l["ts"], l["signatur"], l["wallet"]))

    # Token-beholdning ved periodens slut (alle egne wallets samlet, hele historikken)
    beholdning = defaultdict(int)
    for l in alle_linjer:
        if l["mint"] and l["dato_dk"] <= til:
            beholdning[l["mint"]] += l["token"]

    handel = [l for l in periode if l["type"] in (C.KOEB, C.SALG)]
    koeb = [l for l in handel if l["type"] == C.KOEB]
    salg = [l for l in handel if l["type"] == C.SALG]
    creator = [l for l in periode if l["type"] == C.CREATOR_FEE]

    wb = Workbook()

    # ---------- Resumé ----------
    ws = wb.active
    ws.title = "Resumé"
    netvaerk_sol, netvaerk_dkk = _sum(periode, "netvaerksgebyr_sol"), _sum(periode, "netvaerksgebyr_dkk")
    jito_sol, jito_dkk = _sum(periode, "jito_sol"), _sum(periode, "jito_dkk")
    handel_sol, handel_dkk = _sum(handel, "sol_sol"), _sum(handel, "sol_dkk")
    creator_sol, creator_dkk = _sum(creator, "sol_sol"), _sum(creator, "sol_dkk")
    resultat_sol = handel_sol + creator_sol - netvaerk_sol - jito_sol
    resultat_dkk = handel_dkk + creator_dkk - netvaerk_dkk - jito_dkk
    mints = {l["mint"] for l in handel}

    raekker = [
        ("Regnskab – memecoin-handel på Solana", None, None),
        (f"Periode: {fra.isoformat()} til {til.isoformat()} (dansk tid)", None, None),
        (f"Antal wallets: {len({l['wallet'] for l in periode})}   Antal tokens handlet: {len(mints)}   "
         f"Heraf egne launches: {len(mints & oprettet)}", None, None),
        (None, None, None),
        ("Post", "SOL", "DKK"),
        ("Salg i alt (modtaget, efter Proxima-salgsgebyr)", _sum(salg, "sol_sol"), _sum(salg, "sol_dkk")),
        ("Køb i alt (betalt, inkl. Proxima-købsgebyr)", _sum(koeb, "sol_sol"), _sum(koeb, "sol_dkk")),
        ("Handelsresultat (salg − køb)", handel_sol, handel_dkk),
        ("Creator fees (indtægt)", creator_sol, creator_dkk),
        ("Netværksgebyrer (udgift)", -netvaerk_sol, -netvaerk_dkk),
        ("Jito tips (udgift)", -jito_sol, -jito_dkk),
        ("RESULTAT", resultat_sol, resultat_dkk),
        (None, None, None),
        ("Til information – allerede indeholdt i køb/salg ovenfor:", None, None),
        ("Proxima købsgebyr", _sum(koeb, "proxima_gebyr_sol"), _sum(koeb, "proxima_gebyr_dkk")),
        ("Proxima salgsgebyr", _sum(salg, "proxima_gebyr_sol"), _sum(salg, "proxima_gebyr_dkk")),
        ("Kontodepositum (rent) netto betalt – kan hentes retur, ikke en udgift",
         _sum(periode, "rent_sol"), _sum(periode, "rent_dkk")),
        (None, None, None),
        ("Bevægelser der IKKE er resultat (se fanen 'Til gennemgang'):", None, None),
        ("SOL ind fra eksterne adresser (indskud?)",
         _sum([l for l in periode if l["type"] == C.SOL_IND], "sol_sol"),
         _sum([l for l in periode if l["type"] == C.SOL_IND], "sol_dkk")),
        ("SOL ud til eksterne adresser (hævning/betaling?)",
         _sum([l for l in periode if l["type"] == C.SOL_UD], "sol_sol"),
         _sum([l for l in periode if l["type"] == C.SOL_UD], "sol_dkk")),
        (f"Interne overførsler mellem egne wallets: {sum(1 for l in periode if l['type'] in (C.INTERN_SOL, C.INTERN_TOKEN))} linjer (udlignes, ingen resultatpåvirkning)", None, None),
    ]
    estimeret = any(l["gebyr_estimeret"] for l in handel)
    if estimeret:
        raekker.append((f"Proxima-gebyrer er ESTIMERET ud fra {cfg['proxima']['koebsgebyr']:.0%} køb / "
                        f"{cfg['proxima']['salgsgebyr']:.0%} salg. Angiv Proximas gebyr-wallet i config for præcise tal.",
                        None, None))
    if manglende:
        raekker.append((f"ADVARSEL: mangler SOL-kurs for {len(manglende)} dage – DKK er ufuldstændig.", None, None))
    for r in raekker:
        ws.append(list(r))
    ws["A1"].font = Font(bold=True, size=14)
    for c in ws[5]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        if row[0].value == "RESULTAT" or row[0].value == "Handelsresultat (salg − køb)":
            for c in row:
                c.font = Font(bold=True)
    _formater(ws, {"B": "#,##0.0000", "C": "#,##0.00"}, bredder={"A": 70, "B": 16, "C": 18}, header_row=None)

    # ---------- Launches ----------
    ws = wb.create_sheet("Launches")
    hdr = ["Token (mint)", "Navn", "Egen launch", "Første handel", "Sidste handel", "Antal wallets",
           "Køb SOL", "Salg SOL", "Netto SOL", "Netto DKK", "Proxima købsgebyr SOL", "Proxima salgsgebyr SOL",
           "Tokens tilbage ved periodens slut", "Solscan"]
    ws.append(hdr)
    per_mint = defaultdict(list)
    for l in handel:
        per_mint[l["mint"]].append(l)
    for m, rows in sorted(per_mint.items(), key=lambda kv: min(r["ts"] for r in kv[1])):
        k = [r for r in rows if r["type"] == C.KOEB]
        s = [r for r in rows if r["type"] == C.SALG]
        rest = beholdning.get(m, 0)
        dec = rows[0]["decimals"]
        ws.append([m, navne.get(m, ""), "Ja" if m in oprettet else "Nej",
                   min(r["tid_dk"] for r in rows), max(r["tid_dk"] for r in rows),
                   len({r["wallet"] for r in rows}),
                   -_sum(k, "sol_sol"), _sum(s, "sol_sol"), _sum(rows, "sol_sol"), _sum(rows, "sol_dkk"),
                   _sum(k, "proxima_gebyr_sol"), _sum(s, "proxima_gebyr_sol"),
                   rest / 10 ** dec if dec else rest, f"https://solscan.io/token/{m}"])
    _formater(ws, {"D": "yyyy-mm-dd hh:mm", "E": "yyyy-mm-dd hh:mm", "G": "#,##0.0000", "H": "#,##0.0000",
                   "I": "#,##0.0000", "J": "#,##0.00", "K": "#,##0.0000", "L": "#,##0.0000", "M": "#,##0.00"},
              bredder={"A": 46, "N": 60})

    # ---------- Wallets ----------
    ws = wb.create_sheet("Wallets")
    ws.append(["Wallet", "Navn", "Antal køb", "Antal salg", "Køb SOL", "Salg SOL", "Handel netto SOL",
               "Creator fees SOL", "Netværksgebyr SOL", "Jito SOL", "Ekstern SOL ind", "Ekstern SOL ud",
               "Resultat DKK", "Solscan"])
    per_w = defaultdict(list)
    for l in periode:
        per_w[l["wallet"]].append(l)
    for w, rows in sorted(per_w.items()):
        k = [r for r in rows if r["type"] == C.KOEB]
        s = [r for r in rows if r["type"] == C.SALG]
        cf = [r for r in rows if r["type"] == C.CREATOR_FEE]
        res_dkk = _sum(k + s + cf, "sol_dkk") - _sum(rows, "netvaerksgebyr_dkk") - _sum(rows, "jito_dkk")
        ws.append([w, navne.get(w, ""), len(k), len(s), -_sum(k, "sol_sol"), _sum(s, "sol_sol"),
                   _sum(k + s, "sol_sol"), _sum(cf, "sol_sol"), _sum(rows, "netvaerksgebyr_sol"),
                   _sum(rows, "jito_sol"),
                   _sum([r for r in rows if r["type"] == C.SOL_IND], "sol_sol"),
                   _sum([r for r in rows if r["type"] == C.SOL_UD], "sol_sol"),
                   res_dkk, f"https://solscan.io/account/{w}"])
    _formater(ws, {c: "#,##0.0000" for c in "EFGHIJKL"} | {"M": "#,##0.00"}, bredder={"A": 46, "N": 60})

    # ---------- Transaktioner ----------
    ws = wb.create_sheet("Transaktioner")
    tx_hdr = ["Tid (dansk)", "Type", "Wallet", "Wallet-navn", "Token (mint)", "Token-navn", "Token-mængde",
              "SOL", "Proxima-gebyr SOL", "Gebyr estimeret", "Netværksgebyr SOL", "Jito SOL", "Rent SOL",
              "SOL/DKK kurs", "DKK", "Signatur", "Solscan"]
    ws.append(tx_hdr)
    for l in periode:
        ws.append(_tx_raekke(l))
    _formater(ws, _tx_formater(), bredder={"C": 46, "E": 46, "P": 30, "Q": 60})

    # ---------- Til gennemgang ----------
    ws = wb.create_sheet("Til gennemgang")
    ws.append(["Hvorfor"] + tx_hdr)
    for l in periode:
        grund = None
        if l["type"] in C.GENNEMGANG:
            grund = "Bevægelse til/fra en adresse uden for dine wallets – indskud, hævning, betaling eller manglende wallet?"
        elif l.get("flere_mints"):
            grund = "Flere tokens i samme transaktion – SOL er fordelt efter mængde"
        elif l["sol_dkk_kurs"] is None:
            grund = "Mangler SOL-kurs"
        if grund:
            ws.append([grund] + _tx_raekke(l))
    _formater(ws, {get_column_letter(ord(k) - 64 + 1): v for k, v in _tx_formater().items()},
              bredder={"A": 60, "D": 46, "F": 46, "R": 60})

    # ---------- Kurser ----------
    ws = wb.create_sheet("Kurser")
    ws.append(["Dato (UTC)", "SOL/USD", "USD/DKK", "SOL/DKK"])
    for d in sorted(kurser):
        if fra <= d <= til:
            ws.append([d, *kurser[d]])
    _formater(ws, {"A": "yyyy-mm-dd", "B": "#,##0.00", "C": "0.0000", "D": "#,##0.00"})

    # ---------- Forklaring ----------
    ws = wb.create_sheet("Forklaring")
    for t in FORKLARING:
        ws.append([t])
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(filnavn)
    return {"resultat_sol": resultat_sol, "resultat_dkk": resultat_dkk, "linjer": len(periode),
            "tokens": len(mints), "egne_launches": len(mints & oprettet), "manglende_kurser": len(manglende)}


def _tx_raekke(l):
    return [l["tid_dk"], l["type"], l["wallet"], l["wallet_navn"], l["mint"] or "", l["navn"],
            l["token_ui"] if l["mint"] else None, l["sol_sol"], l["proxima_gebyr_sol"] or None,
            ("Ja" if l["gebyr_estimeret"] else "Nej") if l["type"] in (C.KOEB, C.SALG) else "",
            l["netvaerksgebyr_sol"] or None, l["jito_sol"] or None, l["rent_sol"] or None,
            l["sol_dkk_kurs"], l["sol_dkk"], l["signatur"], f"https://solscan.io/tx/{l['signatur']}"]


def _tx_formater():
    return {"A": "yyyy-mm-dd hh:mm:ss", "G": "#,##0.00", "H": "#,##0.000000", "I": "#,##0.000000",
            "K": "#,##0.000000", "L": "#,##0.000000", "M": "#,##0.000000", "N": "#,##0.00", "O": "#,##0.00"}


def _formater(ws, talformater, bredder=None, header_row=1):
    if header_row:
        for c in ws[header_row]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="333333")
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        ws.auto_filter.ref = ws.dimensions
    for col, fmt in talformater.items():
        for c in ws[col]:
            if isinstance(c.value, (int, float)) or hasattr(c.value, "year"):
                c.number_format = fmt
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    for col, w in (bredder or {}).items():
        ws.column_dimensions[col].width = w


FORKLARING = [
    "SÅDAN ER TALLENE LAVET",
    "",
    "Kilde: Alle transaktioner er hentet direkte fra Solana-blockchainen for de wallets der står i config-filen. "
    "Hver linje kan efterprøves via Solscan-linket.",
    "",
    "Beløb: SOL-ændringen er beregnet ud fra walletens saldo før og efter hver transaktion. Køb og salg er derfor "
    "de beløb der faktisk gik ud af og ind på walleten – Proximas 5% købsgebyr og 5% salgsgebyr er allerede "
    "indeholdt i dem. Gebyret vises separat i kolonnen 'Proxima-gebyr' til information, så det ikke trækkes to gange.",
    "",
    "Proxima-gebyr: Hvis Proximas gebyr-wallet er angivet i config, er gebyret aflæst præcist på kæden. Ellers er det "
    "estimeret: ved køb = betalt × 5/105, ved salg = modtaget × 5/95.",
    "",
    "Netværksgebyrer og Jito tips er trukket ud af handelsbeløbet og vist som separate udgifter.",
    "",
    "Rent (kontodepositum): Når en wallet første gang modtager et token, betales ca. 0,002 SOL i depositum for "
    "token-kontoen. Det får man retur når kontoen lukkes. Det er trukket ud af handelsbeløbet og er ikke en udgift.",
    "",
    "Interne overførsler: SOL eller tokens flyttet mellem dine egne wallets er hverken indtægt eller udgift.",
    "",
    "Creator fees: SOL modtaget fra pump.fun-programmerne uden et token-salg tolkes som creator fee (indtægt).",
    "",
    "Til gennemgang: Bevægelser til/fra adresser der ikke er på din wallet-liste. Det kan være indskud fra privat, "
    "hævning til en børs, betaling – eller en wallet du har glemt at tilføje. Gennemgå dem med din revisor.",
    "",
    "Kurser: SOL/DKK = SOL/USD (Binance, dagens lukkekurs, UTC) × USD/DKK (ECB-referencekurs, som Danmarks "
    "Nationalbank også bruger). Weekender bruger seneste kendte kurs.",
    "",
    "Usolgte tokens: Kolonnen 'Tokens tilbage' på fanen Launches viser tokens der ikke er solgt ved periodens slut. "
    "Købsprisen for dem indgår i nettotallet. Din revisor skal vurdere hvordan de skal værdiansættes.",
    "",
    "Skat: Opgørelsen er et grundlag for regnskabet, ikke en skatteberegning. Selskabet beskattes med 22%, men "
    "opgørelsesprincip (fx realisation/lager) og behandlingen af selve SOL-beholdningen skal fastlægges med revisor.",
]
